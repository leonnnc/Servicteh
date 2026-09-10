# -*- coding: utf-8 -*-
"""Genera plantilla-servitech.xlsx: estructura de datos para AppSheet (Servitech)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- Tema ----
FONT   = 'Microsoft YaHei'
HEADER = '0E6B66'   # verde teal Servitech
WHITE  = 'FFFFFF'
BAND   = 'EAF4F3'
AMBER  = 'FFF3CD'
BORDER = 'BFCFCD'
GRAY   = '5F6B6A'

thin = Side(style='thin', color=BORDER)
box  = Border(left=thin, right=thin, top=thin, bottom=thin)
hfont  = Font(name=FONT, size=12, bold=True, color=WHITE)
hfill  = PatternFill('solid', start_color=HEADER)
amber  = PatternFill('solid', start_color=AMBER)
center = Alignment(horizontal='center', vertical='center')

def wlen(s):
    return sum(2 if ord(ch) > 127 else 1 for ch in str(s))

def build_table(ws, headers, example, tab, fmts=None, note=None):
    ws.sheet_properties.tabColor = tab
    fmts = fmts or {}
    if note:
        ws.cell(1, 1, note).font = Font(name=FONT, size=11, italic=True, color=GRAY)
        start = 2
    else:
        start = 1
    for j, name in enumerate(headers, 1):
        c = ws.cell(start, j, name)
        c.font = hfont; c.fill = hfill; c.alignment = center; c.border = box
    r = start + 1
    for j, val in enumerate(example, 1):
        c = ws.cell(r, j, val)
        c.border = box; c.fill = amber
        c.alignment = Alignment(horizontal='right' if isinstance(val, (int, float)) else 'left')
        if headers[j-1] in fmts:
            c.number_format = fmts[headers[j-1]]
    for j, name in enumerate(headers, 1):
        col = get_column_letter(j)
        best = wlen(name)
        if example[j-1] is not None:
            best = max(best, wlen(example[j-1]))
        ws.column_dimensions[col].width = min(max(best + 3, 10), 34)

wb = Workbook()
wb._named_styles['Normal'].font = Font(name=FONT, size=11)

# ---------- 0. Instrucciones ----------
ins = wb.active
ins.title = 'Instrucciones'
ins.sheet_properties.tabColor = '5F6B6A'
ins.column_dimensions['A'].width = 14
ins.column_dimensions['B'].width = 96
rows = [
    ('PLANTILLA', 'App de soporte tecnico Servitech - estructura de datos para AppSheet', True),
    ('', '', False),
    ('Que es', 'Este archivo tiene 5 hojas (Empresas, Equipos, Tareas, Repuestos, Informes) con las columnas', False),
    ('', 'que usara tu app. Cada hoja = una tabla. Contiene 1 fila de EJEMPLO (amarilla) para que', False),
    ('', 'AppSheet reconozca los tipos de dato. Borra esas filas cuando la app este creada.', False),
    ('', '', False),
    ('PASO 1', 'Guarda este archivo en tu computadora y subelo a Google Drive (drive.google.com).', False),
    ('PASO 2', 'Abrelo en Drive con Google Sheets (clic derecho > Abrir con > Google Sheets).', False),
    ('PASO 3', 'Entra a appsheet.com con tu misma cuenta de Google y elige "Crear app" > Google Sheets.', False),
    ('PASO 4', 'Selecciona este archivo de Drive: AppSheet detectara las 5 hojas como tablas.', False),
    ('PASO 5', 'Borra la fila de EJEMPLO de cada hoja (dentro de AppSheet o en la hoja de calculo).', False),
    ('PASO 6', 'Instala "AppSheet" desde Google Play en tu Android e inicia sesion con el mismo Google.', False),
    ('', '', False),
    ('RELACIONES', 'Conecta las tablas en AppSheet (columna tipo Ref) para que todo quede enlazado:', True),
    ('Hoja', 'Columna que se conecta', True),
    ('Equipos', 'id_empresa  ->  tabla Empresas', False),
    ('Tareas', 'id_empresa  ->  tabla Empresas', False),
    ('Tareas', 'id_equipo   ->  tabla Equipos (opcional)', False),
    ('Repuestos', 'id_tarea    ->  tabla Tareas', False),
    ('Informes', 'id_tarea    ->  tabla Tareas', False),
    ('', '', False),
    ('TIPOS', 'Como llenar cada tipo de dato:', True),
    ('Si/No', 'Escribir: Si  o  No  (ej. activo)', False),
    ('Fechas', 'Formato: anio-mes-dia, ej. 2026-09-09', False),
    ('Precio', 'Numero con punto decimal, ej. 125.50', False),
    ('Fotos/firma', 'Columnas fotos, firma_responsable, firma_tecnico y pdf las llena la app (imagen/archivo).', False),
    ('', '', False),
    ('IMPORTANTE', 'No borres ni renombres las columnas de los encabezados: rompe la app.', True),
    ('', 'Al terminar la conexion avisame y te guio en el siguiente bloque:', False),
    ('Siguiente', 'Formulario de tarea rapido, campo de firma, generacion del PDF del informe y envio por email/WhatsApp.', False),
]
r = 1
for label, text, head in rows:
    lc = ins.cell(r, 1, label); tc = ins.cell(r, 2, text)
    lc.font = Font(name=FONT, size=11, bold=head, color=WHITE if head else HEADER)
    if head and label != '':
        lc.fill = hfill; lc.alignment = center
    tc.font = Font(name=FONT, size=11, bold=head, color=HEADER if head else '333333')
    lc.alignment = Alignment(vertical='center')
    tc.alignment = Alignment(vertical='center')
    r += 1

# ---------- 1. Empresas ----------
ws = wb.create_sheet('Empresas')
emp = ['id_empresa', 'razon_social', 'ruc', 'direccion', 'telefono', 'email',
       'persona_contacto', 'cargo_contacto', 'telefono_contacto', 'rubro',
       'notas', 'fecha_alta', 'activo']
ex  = ['99999 (EJEMPLO)', 'EMPRESA DEMO S.A.C. (ejemplo)', '20123456789', 'Av. Principal 123, Lima',
       '01 555 1234', 'contacto@demo.com', 'Juan Perez', 'Administrador', '999 888 777',
       'Oficina', 'Borrar esta fila despues de crear la app', '2026-09-09', 'Si']
build_table(ws, emp, ex, '0E6B66',
            fmts={'fecha_alta': 'yyyy-mm-dd'})

# ---------- 2. Equipos ----------
ws = wb.create_sheet('Equipos')
eq = ['id_equipo', 'id_empresa', 'tipo_equipo', 'marca', 'modelo', 'nro_serie',
      'ubicacion', 'fecha_instalacion', 'garantia_hasta', 'notas_equipo', 'activo']
ex = ['9999 (EJEMPLO)', '99999 (EJEMPLO)', 'UPS', 'APC', 'Smart-UPS 1500', 'AS1234567890',
      'Piso 2 - Sala servidores', '2025-03-15', '2027-03-15', 'Borrar esta fila', 'Si']
build_table(ws, eq, ex, '2E6FA3',
            fmts={'fecha_instalacion': 'yyyy-mm-dd', 'garantia_hasta': 'yyyy-mm-dd'})

# ---------- 3. Tareas ----------
ws = wb.create_sheet('Tareas')
ta = ['id_tarea', 'id_empresa', 'id_equipo', 'tipo_tarea', 'prioridad', 'estado',
      'descripcion_trabajo', 'novedad', 'trabajo_realizado', 'solucion', 'equipo_operativo',
      'recomendaciones', 'tecnico_responsable', 'fecha_creacion', 'fecha_programada',
      'fecha_inicio', 'fecha_fin', 'informe_emitido', 'fotos']
ex = ['T-9999 (EJEMPLO)', '99999 (EJEMPLO)', '9999 (EJEMPLO)', 'Correctivo', 'Alta', 'Pendiente',
      'UPS no enciende, reporte del cliente', 'Fuente de poder danada', '', '', '',
      'Revisar baterias', 'Tu Nombre', '2026-09-09', '2026-09-10', '', '', 'No', '']
build_table(ws, ta, ex, 'C56443',
            fmts={'fecha_creacion': 'yyyy-mm-dd', 'fecha_programada': 'yyyy-mm-dd'})

# ---------- 4. Repuestos ----------
ws = wb.create_sheet('Repuestos')
rp = ['id_repuesto', 'id_tarea', 'id_equipo', 'descripcion_pieza', 'referencia',
      'cantidad', 'precio_unitario', 'proveedor', 'estado_pedido',
      'fecha_pedido', 'fecha_llegada', 'notas_repuesto']
ex = ['R-9999 (EJEMPLO)', 'T-9999 (EJEMPLO)', '9999 (EJEMPLO)', 'Fuente de poder 500W',
      'FP-500', 1, 189.50, 'Casa de baterias', 'Por comprar', '', '', 'Borrar esta fila']
build_table(ws, rp, ex, '7B5AA6',
            fmts={'cantidad': '#,##0', 'precio_unitario': '#,##0.00'})

# ---------- 5. Informes ----------
ws = wb.create_sheet('Informes')
inf = ['id_informe', 'id_tarea', 'fecha_emision', 'tecnico',
       'emp_razon_social', 'emp_ruc', 'emp_direccion', 'emp_contacto',
       'eq_tipo', 'eq_marca', 'eq_modelo', 'eq_serie',
       'novedad', 'trabajo_realizado', 'repuestos_cambiados', 'solucion',
       'nombre_responsable', 'cargo_responsable', 'firma_responsable', 'firma_tecnico',
       'enviado_a', 'fecha_envio', 'pdf']
ex = ['0001-2026 (EJEMPLO)', 'T-9999 (EJEMPLO)', '2026-09-09', 'Tu Nombre',
      'EMPRESA DEMO S.A.C. (ejemplo)', '20123456789', 'Av. Principal 123, Lima', 'Juan Perez',
      'UPS', 'APC', 'Smart-UPS 1500', 'AS1234567890',
      'Fuente de poder danada', 'Se reemplazo la fuente de poder', '1x Fuente de poder 500W - 189.50', 'Equipo operativo',
      '', '', '', '', '', '', '']
build_table(ws, inf, ex, '4E8A4E',
            fmts={'fecha_emision': 'yyyy-mm-dd'})

wb.save('plantilla-servitech.xlsx')
print('OK: plantilla-servitech.xlsx generada')
